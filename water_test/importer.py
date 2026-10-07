import csv
import io
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook

from .models import LabTest

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
TECHNICIANS = [name for name, _ in LabTest.TECHNICIAN_CHOICES]

# header (lower-cased, trimmed) -> field
HEADER_ALIASES = {
    "date": "date", "date_created": "date", "date created": "date", "date/time": "date", "datetime": "date",
    "sample": "sample_name", "sample name": "sample_name", "sample_name": "sample_name",
    "technician": "technician", "tech": "technician",
    "tds": "tds", "ec": "ec", "conductivity": "ec", "ph": "ph", "salinity": "salinity",
    "temp": "temperature", "temperature": "temperature",
}
REQUIRED = ["date", "sample_name", "technician", "tds", "ec", "ph", "salinity", "temperature"]
DECIMAL_FIELDS = ["tds", "ec", "ph", "salinity", "temperature"]

# Only physically impossible values are rejected — these are not the standard's
# acceptance limits, so a genuinely out-of-spec reading is still recorded as measured.
PLAUSIBLE = {"tds": (0, 100000), "ec": (0, 100000), "ph": (0, 14), "salinity": (0, 1000), "temperature": (0, 60)}
LABELS = {"tds": "TDS", "ec": "EC", "ph": "pH", "salinity": "Salinity", "temperature": "Temp"}

DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
    "%d-%m-%Y %H:%M", "%d-%m-%Y",
)


def _read_table(upload):
    if upload.size > MAX_BYTES:
        raise ValueError("The file is larger than 5 MB.")
    name = upload.name.lower()
    data = upload.read()
    if name.endswith(".xlsx"):
        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        except Exception:
            raise ValueError("That doesn't look like a valid .xlsx file.")
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")
        return list(csv.reader(io.StringIO(text)))
    raise ValueError("Upload an .xlsx or .csv file.")


def _parse_date(value):
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, str):
        text = value.strip()
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
    return None


def _parse_decimal(value):
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip()
    if "," in text:
        # "9,5" could be a decimal comma or "95" with a stray separator — guessing would
        # silently change a measured value, so make the user fix the source instead.
        return None
    if not text:
        return None
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None
    if not number.is_finite():
        return None
    return number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _existing_keys(dates):
    if not dates:
        return set()
    rows = LabTest.objects.filter(
        date_created__date__gte=min(dates), date_created__date__lte=max(dates)
    ).values_list("sample_name", "date_created")
    return {(name, stamp.date()) for name, stamp in rows}


def parse_upload(upload):
    """Reads an uploaded .xlsx/.csv and sorts every row into: `rows` (ready to import),
    `duplicates` (already in Lab Test — same sample name on the same day, left untouched)
    and `errors` (can't be imported, with the reason). Raises ValueError for a
    problem with the file as a whole. Values are returned as strings so the result can
    sit in the session between the preview and the confirmation."""
    table = _read_table(upload)
    header_index = next((i for i, row in enumerate(table) if any(not _blank(c) for c in row)), None)
    if header_index is None:
        raise ValueError("The file is empty.")

    columns = {}
    for position, cell in enumerate(table[header_index]):
        field = HEADER_ALIASES.get(str(cell).strip().lower()) if cell is not None else None
        if field and field not in columns:
            columns[field] = position
    missing = [LABELS.get(f, f.replace("_", " ").title()) for f in REQUIRED if f not in columns]
    if missing:
        raise ValueError(f"Missing column(s): {', '.join(missing)}. Expected: Date, Sample, Technician, TDS, EC, pH, Salinity, Temp.")

    body = [(header_index + 2 + i, row) for i, row in enumerate(table[header_index + 1:])]
    body = [(n, row) for n, row in body if any(not _blank(c) for c in row)]
    if len(body) > MAX_ROWS:
        raise ValueError(f"The file has {len(body)} rows; the limit is {MAX_ROWS} per import.")
    if not body:
        raise ValueError("The file has a header but no data rows.")

    technicians = {t.lower(): t for t in TECHNICIANS}
    today = datetime.now()
    parsed, errors = [], []
    for number, row in body:
        def cell(field):
            index = columns[field]
            return row[index] if index < len(row) else None

        problems = []
        stamp = _parse_date(cell("date"))
        if stamp is None:
            problems.append("Date is missing or not a valid date")
        elif stamp > today + timedelta(days=1):
            problems.append("Date is in the future")
        elif stamp.year < 2000:
            problems.append("Date is before 2000")

        sample = "" if _blank(cell("sample_name")) else str(cell("sample_name")).strip()
        if not sample:
            problems.append("Sample name is missing")
        elif len(sample) > 255:
            problems.append("Sample name is too long")

        raw_tech = "" if _blank(cell("technician")) else str(cell("technician")).strip()
        technician = technicians.get(raw_tech.lower())
        if technician is None:
            problems.append(
                f"Technician '{raw_tech}' is not recognised (use one of: {', '.join(TECHNICIANS)})" if raw_tech
                else "Technician is missing"
            )

        values = {}
        for field in DECIMAL_FIELDS:
            number_value = _parse_decimal(cell(field))
            low, high = PLAUSIBLE[field]
            if number_value is None:
                raw = cell(field)
                if isinstance(raw, str) and "," in raw:
                    problems.append(f"{LABELS[field]} '{raw}' contains a comma — use a decimal point")
                else:
                    problems.append(f"{LABELS[field]} is missing or not a number")
            elif not low <= number_value <= high:
                problems.append(f"{LABELS[field]} {number_value} is outside the possible range {low}–{high}")
            else:
                values[field] = number_value

        if problems:
            errors.append({"row": number, "message": "; ".join(problems)})
            continue
        parsed.append({
            "row": number, "sample_name": sample, "technician": technician,
            "date_created": stamp.isoformat(timespec="seconds"),
            **{field: str(value) for field, value in values.items()},
        })

    existing = _existing_keys([datetime.fromisoformat(r["date_created"]).date() for r in parsed])
    seen, rows, duplicates = {}, [], []
    for record in parsed:
        key = (record["sample_name"], datetime.fromisoformat(record["date_created"]).date())
        if key in seen:
            errors.append({"row": record["row"], "message": f"Same sample and date as row {seen[key]} in this file"})
        elif key in existing:
            duplicates.append(record)
            seen[key] = record["row"]
        else:
            rows.append(record)
            seen[key] = record["row"]
    errors.sort(key=lambda e: e["row"])
    return {"rows": rows, "duplicates": duplicates, "errors": errors, "total": len(body)}


@transaction.atomic
def save_rows(rows):
    """Writes previewed rows, re-checking duplicates first — someone may have added the
    same result between the preview and the confirmation. Returns (created, skipped)."""
    dates = [datetime.fromisoformat(r["date_created"]).date() for r in rows]
    existing = _existing_keys(dates)
    to_create, skipped = [], 0
    for r in rows:
        stamp = datetime.fromisoformat(r["date_created"])
        if (r["sample_name"], stamp.date()) in existing:
            skipped += 1
            continue
        to_create.append(LabTest(
            sample_name=r["sample_name"], technician=r["technician"], date_created=stamp,
            tds=Decimal(r["tds"]), ec=Decimal(r["ec"]), ph=Decimal(r["ph"]),
            salinity=Decimal(r["salinity"]), temperature=Decimal(r["temperature"]),
        ))
    LabTest.objects.bulk_create(to_create)
    return len(to_create), skipped
